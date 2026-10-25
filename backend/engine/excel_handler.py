import re
from engine.mapping_engine import MappingEngine
from rules import RedundancyNormalizer, HardwareSystemConfig
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from typing import List, Dict, Any, Tuple, Optional
from models import (
    SignalRecord, AssignmentModel, ValidationIssueModel,
    TQResponseModel, UtilisationSummaryModel, RevisionComparisonResponse
)

if os.environ.get("VERCEL"):
    REPORT_OUTPUT_DIR = "/tmp/reports"
else:
    REPORT_OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "..", "data", "reports")

try:
    os.makedirs(REPORT_OUTPUT_DIR, exist_ok=True)
except Exception:
    REPORT_OUTPUT_DIR = "/tmp/reports"
    os.makedirs(REPORT_OUTPUT_DIR, exist_ok=True)

class ExcelHandler:
    @staticmethod
    def inspect_workbook(file_path: str) -> Dict[str, Any]:
        """
        Reads workbook sheets (Excel or CSV), headers, sample rows and computes fuzzy column mappings.
        """
        is_csv = file_path.lower().endswith('.csv')
        previews = {}
        headers_map = {}
        suggested_map = {}
        sheet_names = []

        if is_csv:
            sheet_names = ["Master_IO_Database"]
            df = None
            for enc in ['utf-8', 'latin1', 'cp1252', 'utf-8-sig']:
                try:
                    df = pd.read_csv(file_path, encoding=enc, nrows=20)
                    break
                except Exception:
                    continue
            
            if df is None:
                raise ValueError("Could not parse CSV file with supported encodings (utf-8, latin1, cp1252).")

            header_row = [str(c).strip() for c in df.columns]
            headers_map["Master_IO_Database"] = header_row

            preview_rows = []
            for _, r in df.head(10).iterrows():
                row_dict = {}
                for h in header_row:
                    val = r.get(h)
                    row_dict[h] = "" if pd.isna(val) else str(val).strip()
                preview_rows.append(row_dict)
            previews["Master_IO_Database"] = preview_rows
            ExcelHandler._compute_suggested_mappings(header_row, suggested_map)
        else:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames

            for sheet in sheet_names:
                ws = wb[sheet]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                header_row = None
                header_row_idx = 0
                for idx, r in enumerate(rows[:10]):
                    non_empty = [c for c in r if c is not None and str(c).strip()]
                    if len(non_empty) >= 2:
                        header_row = [str(c).strip() if c is not None else f"COL_{i+1}" for i, c in enumerate(r)]
                        header_row_idx = idx
                        break

                if not header_row:
                    continue

                headers_map[sheet] = header_row

                preview_rows = []
                for r in rows[header_row_idx + 1 : header_row_idx + 11]:
                    row_dict = {}
                    for h, val in zip(header_row, r):
                        row_dict[h] = str(val) if val is not None else ""
                    preview_rows.append(row_dict)

                previews[sheet] = preview_rows
                ExcelHandler._compute_suggested_mappings(header_row, suggested_map)

            wb.close()

        return {
            "sheets": sheet_names,
            "headers": headers_map,
            "previews": previews,
            "suggested_mappings": suggested_map
        }

    @staticmethod
    def _compute_suggested_mappings(header_row: List[str], suggested_map: Dict[str, str], details_map: Optional[Dict[str, Any]] = None):
        scored = MappingEngine.compute_scored_mappings(header_row)
        for field, info in scored.items():
            suggested_map[field] = info["header"]
            if details_map is not None:
                details_map[field] = info
        norm_headers = [(h, h.lower().replace("_", " ").replace("-", " ").strip()) for h in header_row]

        priorities = [
            (
                "tag",
                ["tag number", "tag no", "tag no.", "instrument tag", "signal tag", "tag_number", "tagnumber", "tag name", "tag"],
                lambda hl: "tag" in hl and not any(x in hl for x in ["old", "service", "cable", "status", "ferrule", "system"])
            ),
            (
                "description",
                ["tag service 1 and 2", "service description", "signal description", "tag service", "service", "signal desc", "io description", "description"],
                lambda hl: any(w in hl for w in ["service", "description", "desc"]) and not any(x in hl for x in ["type", "equipment"])
            ),
            (
                "signal_type",
                ["i/o type name", "io type name", "i/o type", "io type", "signal type", "i/o type name 2", "point type"],
                lambda hl: any(hl == w for w in ["signal type", "io type"]) or ("i/o" in hl and "type" in hl)
            ),
            (
                "cable_id",
                ["multipair cable no.", "multipair cable no", "multi cable / cable", "cable tag", "field cable tag", "cable id", "cable no", "cable number"],
                lambda hl: "cable" in hl and any(x in hl for x in ["no", "tag", "num", "id"]) and not any(x in hl for x in ["term", "length", "type"])
            ),
            (
                "pair_number",
                ["pair number", "pair no.", "pair no", "field cable pair", "wire pair", "pair_number", "pair"],
                lambda hl: "pair" in hl and "multipair" not in hl
            ),
            (
                "redundancy",
                ["io grouping", "segregation", "redundancy designator", "redundancy", "system 1 (system)", "system", "voting"],
                lambda hl: any(x in hl for x in ["redund", "grouping", "segreg"])
            ),
            (
                "logical_group",
                ["process function type", "operating area", "design area", "unit", "logical group", "loop name"],
                lambda hl: any(x in hl for x in ["area", "unit", "loop name", "process function"])
            ),
        ]

        for field, exact_candidates, _ in priorities:
            if field in suggested_map:
                continue
            for cand in exact_candidates:
                cand_norm = cand.lower().replace("_", " ").replace("-", " ").strip()
                for original_h, hl in norm_headers:
                    if hl == cand_norm:
                        suggested_map[field] = original_h
                        break
                if field in suggested_map:
                    break

        for field, _, fallback_fn in priorities:
            if field in suggested_map:
                continue
            for original_h, hl in norm_headers:
                if fallback_fn(hl):
                    suggested_map[field] = original_h
                    break

    @staticmethod
    def parse_mapped_signals(file_path: str, sheet_name: str, mapping: Dict[str, str]) -> List[SignalRecord]:
        """
        Parses imported signals based on user-approved column mappings from CSV or Excel.
        """
        is_csv = file_path.lower().endswith('.csv')
        df = None

        if is_csv:
            for enc in ['utf-8', 'latin1', 'cp1252', 'utf-8-sig']:
                try:
                    df = pd.read_csv(file_path, encoding=enc)
                    break
                except Exception:
                    continue
        else:
            df = pd.read_excel(file_path, sheet_name=sheet_name)

        if df is None:
            raise ValueError(f"Could not read source dataset from {file_path}")

        signals: List[SignalRecord] = []

        tag_col = mapping.get("tag")
        type_col = mapping.get("signal_type")
        desc_col = mapping.get("description")
        cable_col = mapping.get("cable_id")
        pair_col = mapping.get("pair_number")
        red_col = mapping.get("redundancy")
        group_col = mapping.get("logical_group")

        if not tag_col:
            raise ValueError("Tag column mapping is required")
        if not type_col:
            raise ValueError("Signal Type column mapping is required")

        for idx, row in df.iterrows():
            if row.isna().all():
                continue

            tag_val = str(row[tag_col]).strip() if tag_col in row and pd.notna(row[tag_col]) else ""
            if tag_val.lower() == "nan":
                tag_val = ""

            type_val = str(row[type_col]).strip().upper() if type_col in row and pd.notna(row[type_col]) else "UNKNOWN"
            if type_val.lower() == "nan":
                type_val = "UNKNOWN"

            desc_val = str(row[desc_col]).strip() if desc_col and desc_col in row and pd.notna(row[desc_col]) else ""
            cable_val = str(row[cable_col]).strip() if cable_col and cable_col in row and pd.notna(row[cable_col]) else ""

            # If all mapped columns are empty, skip row
            if not tag_val and type_val == "UNKNOWN" and not desc_val and not cable_val and (not pair_col or pair_col not in row or pd.isna(row[pair_col])):
                continue

            pair_val = None
            pair_invalid = False
            if pair_col and pair_col in row and pd.notna(row[pair_col]):
                pair_text = str(row[pair_col]).strip()
                if not pair_text or pair_text.upper() in ["-", "N/A", "NA", "NONE", "NULL", "INTERNAL WIRING", "INTERNAL-WIRING"]:
                    pair_val = None
                    pair_invalid = False
                else:
                    match = re.fullmatch(r"(?:pair\s*|p|triad\s*)?(\d+)", pair_text, flags=re.IGNORECASE)
                    if match:
                        pair_val = int(match.group(1))
                        if pair_val < 1:
                            pair_val = None
                            pair_invalid = True
                    else:
                        pair_invalid = True

            _cfg = HardwareSystemConfig()
            raw_red = str(row[red_col]).strip() if red_col and red_col in row and pd.notna(row[red_col]) else ""
            red_val = RedundancyNormalizer.normalize(raw_red, _cfg.redundancy_mappings) or ""

            group_val = str(row[group_col]).strip() if group_col and group_col in row and pd.notna(row[group_col]) else ""

            # Check validity
            is_valid = True
            errors = []
            if not tag_val:
                is_valid = False
                errors.append("Empty Tag name")
            if type_val == "UNKNOWN":
                is_valid = False
                errors.append("Missing signal type")
            if pair_invalid:
                is_valid = False
                errors.append(f"Malformed wire pair value: {row[pair_col]!r}")

            # Check for extra DeltaV metadata in customer columns
            inst_type = ""
            for c in ["Instrument Type", "Instrument", "Type"]:
                if c in row and pd.notna(row[c]):
                    inst_type = str(row[c]).strip()
                    break

            sig_level = ""
            for c in ["Signal Level", "Level", "Wire Type"]:
                if c in row and pd.notna(row[c]):
                    sig_level = str(row[c]).strip()
                    break

            elec_spec = ""
            for c in ["Signal Type", "Electrical Spec", "Power Specification"]:
                if c in row and pd.notna(row[c]):
                    elec_spec = str(row[c]).strip()
                    break

            ext_pwr = ""
            for c in ["External Power Supply", "Power Supply", "External Power"]:
                if c in row and pd.notna(row[c]):
                    ext_pwr = str(row[c]).strip()
                    break

            loop_nm = ""
            for c in ["Loop Name", "Loop", "Internal Loop Order"]:
                if c in row and pd.notna(row[c]):
                    loop_nm = str(row[c]).strip()
                    break

            tech_room = ""
            for c in ["Technical Room (Signal Location 2)", "Technical Room", "Room", "Location"]:
                if c in row and pd.notna(row[c]):
                    tech_room = str(row[c]).strip()
                    break

            op_area = ""
            for c in ["Design Area", "Operating Area", "Area", "Plant Area"]:
                if c in row and pd.notna(row[c]):
                    op_area = str(row[c]).strip()
                    break

            status_val = ""
            for c in ["Status", "Tag Status", "Signal Status"]:
                if c in row and pd.notna(row[c]):
                    status_val = str(row[c]).strip()
                    break

            signals.append(SignalRecord(
                id=idx + 1,
                tag=tag_val,
                description=desc_val,
                signal_type=type_val,
                signal_level=sig_level,
                electrical_spec=elec_spec,
                instrument_type=inst_type,
                external_power=ext_pwr,
                cable_id=cable_val,
                pair_number=pair_val,
                redundancy=red_val,
                logical_group=group_val,
                loop_name=loop_nm,
                technical_room=tech_room,
                operating_area=op_area,
                tag_status=status_val if status_val else None,
                is_valid=is_valid,
                validation_errors=errors if errors else None
            ))

        return signals

    # Styling helpers for OpenPyXL
    @staticmethod
    def _apply_header_style(ws, cols_count: int, title: str = "IO SIGNAL SUDOKU REPORT"):
        navy_fill = PatternFill(start_color="0B132B", end_color="0B132B", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
        sub_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1C2541", end_color="1C2541", fill_type="solid")

        # Title row
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols_count)
        top_cell = ws.cell(row=1, column=1)
        top_cell.value = f"EMERSON DeltaV™ HARDWARE ALLOCATION PLATFORM — {title.upper()}"
        top_cell.font = title_font
        top_cell.fill = navy_fill
        top_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Header row 2
        for col in range(1, cols_count + 1):
            c = ws.cell(row=2, column=col)
            c.font = sub_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 24

    @staticmethod
    def _auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len and "\n" not in val:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # --- REPORT GENERATORS ---

    @classmethod
    def export_completed_allocation(
        cls,
        project_name: str,
        revision_code: str,
        assignments: List[AssignmentModel],
        signals: Optional[List[SignalRecord]] = None
    ) -> str:
        filename = f"{project_name.replace(' ', '_')}_{revision_code}_Completed_Allocation.xlsx"
        file_path = os.path.join(REPORT_OUTPUT_DIR, filename)

        template_path = os.path.join(os.path.dirname(__file__), "..", "sample_files", "official", "Output_Allocation_Template.xlsx")
        
        headers_76 = [
            "Sr. No.", "Revision No", "Tag Name", "Technical Room", "Operating Area",
            "Risk Area", "Unit", "Segregation", "DSTName", "Service Description",
            "Location", "IO Type", "Signal Type", "Signal Level", "NON-IS / IS",
            "Equipment Name", "External Power", "System Cabinet", "Controller",
            "Electronic/ Marshalling Cabinet", "CIOC/CSLS Name", "Baseplate",
            "Channel_No", "CHARM_No", "CHARM Type", "Charm Description",
            "CHARM TB Part No", "CHARM TB Description", "CHARM TB1", "CHARM TB2",
            "CHARM TB3", "CHARM Baseplate Adapter", "M_FTA", "M_FTA Part No",
            "M_FTA_TB", "M_FTA_TB1", "M_FTA_TB2", "M_FTA_TB3", "M_TB", "M_TB1",
            "M_TB2", "M_TB3", "IRC / AUX Cabinet", "IRC_TB", "IRC_TB1", "IRC_TB2",
            "Relay / Splitter", "Relay Coil A1 / Splitter In+", "Relay Coil A2 / Splitter In-",
            "Relay Contact (+)", "Relay Contact (-)", "Field Cable No", "Field Cable Size",
            "Field Cable Pair", "IRC_FTB", "IRC_FTB1", "IRC_FTB2", "Elec Cable No",
            "Elec Cable Size", "Elec Cable Pair", "Tag Status", "Remarks", "BATCH",
            "HW Loop Typical", "System 2", "Old Tag", "Instrument type", "FTA_Cable Name",
            "FTA_Cable Ferrule Name", "Cabinet Side (F/R)", "Change Type",
            "Cabinet Name for Visualization", "Splitter/Relay/\nExternal PWR\nTB(+)",
            "Splitter/Relay/\nExternal PWR\nTB(-)", "Loop Name", "Section / Sub Unit"
        ]

        if os.path.exists(template_path):
            wb = openpyxl.load_workbook(template_path)
            ws = wb["IO Assignment"] if "IO Assignment" in wb.sheetnames else wb.active
            # Remove any pre-existing data rows in template beyond row 1
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "IO Assignment"
            for col_idx, h in enumerate(headers_76, 1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.font = Font(name="Calibri", bold=True)

        sig_map: Dict[str, SignalRecord] = {}
        if signals:
            for s in signals:
                if s.tag:
                    sig_map[s.tag.strip().upper()] = s

        hw_assignments = [a for a in assignments if a.status != "Deleted" and getattr(a, "tag_status", "") != "Deleted"]
        for row_idx, a in enumerate(hw_assignments, 2):
            s = sig_map.get(a.tag.strip().upper()) if a.tag else None

            # Base identifiers and status determination
            sr_no = getattr(a, "sr_no", None) or (row_idx - 1)
            rev_no = getattr(a, "revision_no", None) or revision_code
            tag_st = getattr(a, "tag_status", None) or ("Field" if a.status == "Allocated" else a.status)
            tag_nm = (a.tag or "").strip()
            if not tag_nm:
                tag_nm = "-" if tag_st in ["Spare Slot", "Wired Spare"] else ""

            # Extract hardware index positions
            bp_val = getattr(a, "baseplate_id", None) or "001-CBP-1"
            bp_m = re.search(r'(\d+)$', bp_val)
            bp_num = int(bp_m.group(1)) if bp_m else 1

            ch_num = getattr(a, "channel_no", None) or a.charm_index or 1
            cioc_val = getattr(a, "cioc_id", None) or "CIOC-01"
            cioc_m = re.search(r'(\d+)$', cioc_val)
            cioc_num = int(cioc_m.group(1)) if cioc_m else 1

            # CHARM sequential position on carrier (1..96)
            charm_num = (bp_num - 1) * 12 + ch_num

            # Elevation: Front (CBP 1..4) vs Rear (CBP 5..8)
            cab_side = "F" if bp_num <= 4 else "R"

            # Technical Room & Operating Unit
            unit_val = getattr(a, "unit", None)
            if not unit_val or unit_val == "-":
                tag_u_match = re.match(r"^(\d{3})", tag_nm) if tag_nm and tag_nm != "-" else None
                if tag_u_match:
                    unit_val = tag_u_match.group(1)
                elif "500" in (a.system_cabinet or "") or "P24" in (a.system_cabinet or "") or "500" in (a.electronic_marshalling_cabinet or ""):
                    unit_val = "500"
                else:
                    unit_val = "320"

            tech_rm = getattr(a, "technical_room", None)
            if not tech_rm or tech_rm == "-":
                tech_rm = "PIB-24" if unit_val == "500" else "PIB-30"

            op_area = getattr(a, "operating_area", None)
            if not op_area or op_area == "-" or (unit_val == "320" and "RA" in str(op_area)):
                op_area = "RA 500-02" if unit_val == "500" else "OA-11"

            risk_area = getattr(a, "risk_area", None)
            if not risk_area or risk_area == "-":
                if unit_val == "500":
                    risk_area = "RA 500-02"
                elif unit_val == "440":
                    risk_area = "RA 440-01"
                else:
                    risk_area = "RA 320-01"

            # Cabinets & Controllers (Emerson DeltaV standard architecture)
            if unit_val == "500":
                sys_cab_def = "P24-DCS-50002-01"
                ctrl_def = "P24DPK500201"
                marsh_cab_def = "P24-DCM-50002-01"
                cioc_naming_prefix = "P24DCIOC5002"
            elif cioc_num > 8:
                sys_cab_def = "P30-DCS-32001-02"
                ctrl_def = "P30DPK320102"
                marsh_cab_def = "P30-DCM-32001-02"
                cioc_naming_prefix = "P30DCIOC3201"
            else:
                sys_cab_def = "P30-DCS-32001-01"
                ctrl_def = "P30DPK320101"
                marsh_cab_def = "P30-DCM-32001-01"
                cioc_naming_prefix = "P30DCIOC3201"

            sys_cab = getattr(a, "system_cabinet", None) or a.cabinet_id or sys_cab_def
            if sys_cab == "-":
                sys_cab = sys_cab_def

            ctrl_val = getattr(a, "controller", None) or ctrl_def
            if ctrl_val == "-":
                ctrl_val = ctrl_def

            marsh_cab = getattr(a, "electronic_marshalling_cabinet", None) or marsh_cab_def
            if marsh_cab == "-":
                marsh_cab = marsh_cab_def

            cioc_val = getattr(a, "cioc_id", None) or f"{cioc_naming_prefix}{cioc_num:03d}"
            if not cioc_val.startswith("P24") and not cioc_val.startswith("P30"):
                cioc_val = f"{cioc_naming_prefix}{cioc_num:03d}"

            cioc_suf = f"{cioc_num % 1000:03d}"
            bp_val = f"{cioc_suf}-CBP-{bp_num}"
            cab_viz = f"{marsh_cab}({cab_side})"

            # DSTName
            if tag_nm and tag_nm != "-":
                dst_val = re.sub(r'[^A-Za-z0-9]', '', tag_nm) + "_"
            else:
                dst_val = "_"

            # Service description & Location
            if tag_st == "Cabinet Alarm":
                svc_desc = f"{marsh_cab} TEMPERATURE" if ("DCM" in tag_nm or "0100" in tag_nm) else f"{sys_cab} TEMPERATURE"
                loc_val = tech_rm
            elif tag_st in ["Spare Slot", "Wired Spare"]:
                svc_desc = "-"
                loc_val = "-"
            else:
                svc_desc = getattr(a, "service_description", None) or a.description or (s.description if s else "-") or "-"
                loc_val = getattr(a, "location", None) or "L"

            # IO Type & Signal Categorization
            if tag_st == "Spare Slot":
                io_type = "-"
                sig_type = "-"
                sig_level = "-"
                non_is = "-"
            elif tag_st == "Cabinet Alarm":
                io_type = "AIR"
                sig_type = "RTD"
                sig_level = "3-Wire"
                non_is = "NON-IS"
            else:
                raw_sig = (a.signal_type or (s.signal_type if s else "") or "").upper()
                if "RTD" in raw_sig:
                    io_type = "AIR"
                    sig_type = "RTD"
                    sig_level = "3-Wire"
                elif raw_sig.startswith("AI") or "4-20" in raw_sig:
                    io_type = "AIR"
                    sig_type = "4-20 ma HART"
                    sig_level = "2-Wire"
                elif raw_sig.startswith("AO"):
                    io_type = "AOR"
                    sig_type = "4-20 ma HART"
                    sig_level = "2-Wire"
                elif raw_sig.startswith("DI") or "DRY" in raw_sig:
                    io_type = "DIR"
                    sig_type = "24 VDC Dry Contact"
                    sig_level = "2-Wire"
                elif raw_sig.startswith("DO") or "SOLENOID" in raw_sig or "SOV" in tag_nm:
                    io_type = "DOR"
                    sig_type = "DO 24V DC Isolated" if ("ISOLATED" in raw_sig or "SOV" in tag_nm) else "24 VDC High Side"
                    sig_level = "2-Wire"
                else:
                    io_type = "AIR"
                    sig_type = "4-20 ma HART"
                    sig_level = "2-Wire"
                non_is = "NON-IS"

            # Segregation Train A / Train B
            seg_val = getattr(a, "segregation", None)
            if not seg_val or seg_val == "-":
                if s and s.segregation and s.segregation != "-":
                    seg_val = s.segregation
                elif s and s.redundancy in ["A", "A-CONTROL", "TR-A"]:
                    seg_val = "A-CONTROL"
                elif s and s.redundancy in ["B", "B-CONTROL", "TR-B"]:
                    seg_val = "B-CONTROL"
                elif tag_st in ["Cabinet Alarm", "Spare Slot"]:
                    seg_val = "-"
                elif cioc_num % 2 == 1:
                    seg_val = "A-CONTROL"
                else:
                    seg_val = "B-CONTROL"

            # Equipment Name
            if tag_st == "Cabinet Alarm":
                eq_name = marsh_cab if "DCM" in tag_nm else sys_cab
            elif tag_st in ["Spare Slot", "Wired Spare"]:
                eq_name = "-"
            else:
                eq_name = getattr(a, "equipment_name", None)
                if not eq_name or eq_name == "-":
                    eq_m = re.match(r"^(\d{3})-([A-Z]+)\s*-(\d+)", tag_nm)
                    if eq_m:
                        u, let, num = eq_m.group(1), eq_m.group(2), eq_m.group(3)
                        eq_name = f"{u}-D-{num[:2]}01"
                    else:
                        eq_name = f"{unit_val}-EQ-01"

            ext_pwr = "-"

            # CHARM Hardware Component Mapping & Terminal Pin designations
            if tag_st == "Spare Slot":
                charm_type_part = "SE6103"
                charm_desc = "CHARM PROTECTION COVER"
                charm_tb_part = "SE4501"
                charm_tb_desc = "Standard CHARM Terminal Block"
                tb1 = "-"
                tb2 = "-"
                tb3 = "-"
            elif sig_type == "RTD":
                charm_type_part = "SE4303T03"
                charm_desc = "RTD Input Charm"
                charm_tb_part = "SE4501"
                charm_tb_desc = "Standard CHARM Terminal Block"
                tb1 = "1"
                tb2 = "2"
                tb3 = "4"
            elif io_type == "AIR":
                charm_type_part = "SE4303T01"
                charm_desc = "AI 4-20 MA HART CHARM"
                charm_tb_part = "SE4501"
                charm_tb_desc = "Standard CHARM Terminal Block"
                tb1 = "1"
                tb2 = "2"
                tb3 = "-"
            elif io_type == "AOR":
                charm_type_part = "SE4304T01"
                charm_desc = "AO 4-20 MA HART CHARM"
                charm_tb_part = "SE4501"
                charm_tb_desc = "Standard CHARM Terminal Block"
                tb1 = "1"
                tb2 = "2"
                tb3 = "-"
            elif io_type == "DIR":
                charm_type_part = "SE4301T02"
                charm_desc = "DISCRETE INPUT CHARM; 24 VDC DRY CONTACT: LOW SIDE SENSE"
                charm_tb_part = "SE4501"
                charm_tb_desc = "Standard CHARM Terminal Block"
                tb1 = "1"
                tb2 = "2"
                tb3 = "-"
            elif io_type == "DOR":
                if sig_type == "DO 24V DC Isolated" or "SOV" in tag_nm or "ISOLATED" in str(getattr(a, "terminal_block", "")).upper():
                    charm_type_part = "SE4302T02"
                    charm_desc = "DO 24V DC Isolated"
                    charm_tb_part = "SE4502"
                    charm_tb_desc = "Fused-Injected-Power Terminal Block"
                    tb1 = "-"
                    tb2 = "2"
                    tb3 = "4"
                else:
                    charm_type_part = "SE4302T01"
                    charm_desc = "DISCRETE OUTPUT CHARM; 24 VDC HIGH SIDE"
                    charm_tb_part = "SE4501"
                    charm_tb_desc = "Standard CHARM Terminal Block"
                    tb1 = "1"
                    tb2 = "2"
                    tb3 = "-"
            else:
                charm_type_part = "SE4303T01"
                charm_desc = "AI 4-20 MA HART CHARM"
                charm_tb_part = "SE4501"
                charm_tb_desc = "Standard CHARM Terminal Block"
                tb1 = "1"
                tb2 = "2"
                tb3 = "-"

            bp_adapt = "-"
            m_fta = "-"
            m_fta_p = "-"
            m_fta_tb = "-"
            m_fta_tb1 = "-"
            m_fta_tb2 = "-"
            m_fta_tb3 = "-"
            m_tb = "-"
            m_tb1 = "-"
            m_tb2 = "-"
            m_tb3 = "-"
            irc_cab = "-"
            irc_tb = "-"
            irc_tb1 = "-"
            irc_tb2 = "-"
            rel_spl = "-"
            rel_a1 = "-"
            rel_a2 = "-"
            rel_pos = "-"
            rel_neg = "-"

            # Field Cable Multi-Pair Assignments
            if tag_st == "Cabinet Alarm":
                fcbl_no = "P24-ICE-0001" if unit_val == "500" else "P30-ICE-0001"
                fcbl_sz = "1TX1.3MM2/FR/OS/NIS"
                fcbl_pr = "Triad 1"
            elif tag_st in ["Spare Slot", "Wired Spare"]:
                fcbl_no = "-"
                fcbl_sz = "-"
                fcbl_pr = "-"
            else:
                fcbl_no = getattr(a, "field_cable_no", None) or a.cable_id or (s.cable_id if s else "-") or "-"
                if fcbl_no and fcbl_no != "-":
                    if "CE" in fcbl_no or "24P" in fcbl_no:
                        fcbl_sz = "24PX1MM2/FR/IND&OS/ARM/NIS"
                    elif "12P" in fcbl_no:
                        fcbl_sz = "12PX1MM2/FR/IND&OS/ARM/NIS"
                    elif "6P" in fcbl_no:
                        fcbl_sz = "6PX1MM2/FR/IND&OS/ARM/NIS"
                    elif "CS" in fcbl_no:
                        fcbl_sz = "24PX1MM2/FR/IND&OS/ARM/NIS"
                    else:
                        fcbl_sz = "24PX1MM2/FR/IND&OS/ARM/NIS"
                else:
                    fcbl_sz = "-"

                pair_val = a.pair_number or (s.pair_number if s else None)
                if pair_val:
                    fcbl_pr = f"Pair {pair_val}"
                elif fcbl_no != "-":
                    fcbl_pr = f"Pair {ch_num}"
                else:
                    fcbl_pr = "-"

            irc_ftb = "-"
            irc_ftb1 = "-"
            irc_ftb2 = "-"
            ecbl_no = "-"
            ecbl_sz = "-"
            ecbl_pr = "-"

            # Remarks & Batch
            if tag_st == "Cabinet Alarm":
                remarks = "Internal Cabinet Alarm."
            elif tag_st == "Wired Spare":
                remarks = "Pre-allocated wired spare"
            elif tag_st == "Spare Slot":
                remarks = "Spare Position with SE6103 Cover"
            else:
                remarks = getattr(a, "remarks", None) or a.reason or "-"

            batch_val = "B1"

            # HW Loop Typical
            if tag_st == "Spare Slot":
                hw_loop = "-"
            elif sig_type == "RTD":
                hw_loop = "SA-AMI-530-4EME-000041 - LT-03"
            elif io_type == "AIR":
                hw_loop = "HW-TYP-41 - LT-01"
            elif io_type == "AOR":
                hw_loop = "HW-TYP-41 - LT-04"
            elif io_type == "DIR":
                hw_loop = "SA-AMI-530-4EME-000029 SH.16"
            elif io_type == "DOR":
                hw_loop = "HW-TYP-DO-01"
            else:
                hw_loop = "HW-TYP-41 - LT-01"

            sys2 = "-"
            old_tag = getattr(a, "old_tag", None) or "-"

            # Instrument Type
            if tag_st in ["Spare Slot", "Cabinet Alarm", "Wired Spare"]:
                inst_type = "-"
            else:
                inst_m = re.search(r'[A-Za-z]+', tag_nm)
                code = inst_m.group(0).upper() if inst_m else ""
                if "LT" in code:
                    inst_type = "dP Level Transmitter"
                elif "LV" in code:
                    inst_type = "Control Valve"
                elif "FT" in code:
                    inst_type = "Flow Transmitter"
                elif "FV" in code:
                    inst_type = "Control Valve"
                elif "PT" in code:
                    inst_type = "Pressure Transmitter"
                elif "PV" in code:
                    inst_type = "Pressure Control Valve"
                elif "TT" in code:
                    inst_type = "Temperature Transmitter"
                elif "TV" in code:
                    inst_type = "Temperature Control Valve"
                elif "ZS" in code or "LS" in code:
                    inst_type = "Limit Switch"
                elif "XV" in code or "SOV" in code:
                    inst_type = "Solenoid Valve"
                else:
                    inst_type = "Transmitter / Sensor"

            fta_cbl = "-"
            fta_fer = "-"
            chg_type = getattr(a, "change_type", None) or "-"
            pwr_pos = "-"
            pwr_neg = "-"

            # Loop Name & Section / Sub Unit
            if tag_st in ["Spare Slot", "Cabinet Alarm", "Wired Spare"]:
                loop_nm = "-"
                sec_unit = "-"
            else:
                lp_m = re.match(r"^(\d{3})-([A-Z])([A-Z]*)\s*-(\d+)", tag_nm)
                if lp_m:
                    loop_nm = f"{lp_m.group(1)}-{lp_m.group(2)} -{lp_m.group(4)}"
                else:
                    loop_nm = f"{tag_nm}_LOOP"

                sec_m = re.search(r'(\d+)[A-Z]*$', tag_nm)
                if sec_m:
                    digits = sec_m.group(1)
                    sec_unit = f"{digits[0]}00" if len(digits) >= 2 else "800"
                else:
                    sec_unit = "800"
            row_vals = [
                sr_no, rev_no, tag_nm, tech_rm, op_area, risk_area, unit_val, seg_val, dst_val, svc_desc,
                loc_val, io_type, sig_type, sig_level, non_is, eq_name, ext_pwr, sys_cab, ctrl_val, marsh_cab,
                cioc_val, bp_val, ch_num, charm_num, charm_type_part, charm_desc, charm_tb_part, charm_tb_desc,
                tb1, tb2, tb3, bp_adapt, m_fta, m_fta_p, m_fta_tb, m_fta_tb1, m_fta_tb2, m_fta_tb3,
                m_tb, m_tb1, m_tb2, m_tb3, irc_cab, irc_tb, irc_tb1, irc_tb2, rel_spl, rel_a1, rel_a2,
                rel_pos, rel_neg, fcbl_no, fcbl_sz, fcbl_pr, irc_ftb, irc_ftb1, irc_ftb2, ecbl_no, ecbl_sz,
                ecbl_pr, tag_st, remarks, batch_val, hw_loop, sys2, old_tag, inst_type, fta_cbl, fta_fer,
                cab_side, chg_type, cab_viz, pwr_pos, pwr_neg, loop_nm, sec_unit
            ]

            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Calibri", size=9)

        wb.save(file_path)
        return file_path

    @classmethod
    def export_validation_report(
        cls,
        project_name: str,
        revision_code: str,
        issues: List[ValidationIssueModel]
    ) -> str:
        filename = f"{project_name.replace(' ', '_')}_{revision_code}_Validation_Report.xlsx"
        file_path = os.path.join(REPORT_OUTPUT_DIR, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Validation Issues"

        headers = ["Issue ID", "Severity", "Rule Name", "Affected Tags", "Hardware Location", "Finding Description", "Recommended Engineering Action", "Blocks Approval"]
        cls._apply_header_style(ws, len(headers), f"Validation Audit Log — {revision_code}")

        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=2, column=col_idx, value=h)

        err_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
        wrn_fill = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")
        inf_fill = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")

        for row_idx, issue in enumerate(issues, 3):
            fill = err_fill if issue.severity == "ERROR" else (wrn_fill if issue.severity == "WARNING" else inf_fill)
            row_vals = [
                issue.issue_id, issue.severity, issue.rule_name,
                ", ".join(issue.affected_tags), issue.affected_hardware or "—",
                issue.message, issue.suggested_action, "YES" if issue.blocks_approval else "NO"
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                if col_idx == 2:
                    cell.fill = fill
                    cell.font = Font(name="Calibri", bold=True)
                    cell.alignment = Alignment(horizontal="center")

        cls._auto_fit_columns(ws)
        wb.save(file_path)
        return file_path

    @classmethod
    def export_tq_report(
        cls,
        project_name: str,
        revision_code: str,
        tqs: List[TQResponseModel]
    ) -> str:
        filename = f"{project_name.replace(' ', '_')}_{revision_code}_Technical_Queries.xlsx"
        file_path = os.path.join(REPORT_OUTPUT_DIR, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Technical Queries"

        headers = ["TQ Code", "Status", "Category", "Affected Signals", "Problem Description", "Root Cause Analysis", "Specific Clarification Required", "Suggested Action", "Resolution Notes", "Created At"]
        cls._apply_header_style(ws, len(headers), f"Technical Queries Register — {revision_code}")

        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=2, column=col_idx, value=h)

        for row_idx, tq in enumerate(tqs, 3):
            row_vals = [
                tq.tq_code, tq.status, tq.category, ", ".join(tq.affected_tags),
                tq.description, tq.root_cause, tq.specific_clarification,
                tq.suggested_action, tq.resolution_notes, tq.created_at[:16]
            ]
            for col_idx, val in enumerate(row_vals, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)

        cls._auto_fit_columns(ws)
        wb.save(file_path)
        return file_path

    @classmethod
    def export_utilisation_report(
        cls,
        project_name: str,
        revision_code: str,
        util: UtilisationSummaryModel
    ) -> str:
        filename = f"{project_name.replace(' ', '_')}_{revision_code}_Hardware_Utilisation.xlsx"
        file_path = os.path.join(REPORT_OUTPUT_DIR, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Cabinet & CIOC Utilisation"

        headers = ["Cabinet ID", "CIOC ID", "Total Baseplates", "Occupied Baseplates", "Total CHARMs", "Used CHARMs", "Reserved Spares", "Available Free Slots", "Utilisation %"]
        cls._apply_header_style(ws, len(headers), f"Hardware Utilisation & Spare Audit — {revision_code}")

        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=2, column=col_idx, value=h)

        row_idx = 3
        for cab in util.cabinets:
            for cioc in cab.ciocs:
                row_vals = [
                    cab.cabinet_id, cioc.cioc_id, cioc.total_baseplates, cioc.used_baseplates,
                    cioc.total_slots, cioc.used_slots, cioc.reserved_spare_slots,
                    cioc.free_slots, f"{cioc.utilization_pct}%"
                ]
                for col_idx, val in enumerate(row_vals, 1):
                    ws.cell(row=row_idx, column=col_idx, value=val)
                row_idx += 1

        cls._auto_fit_columns(ws)
        wb.save(file_path)
        return file_path

    @classmethod
    def export_revision_diff(
        cls,
        project_name: str,
        diff: RevisionComparisonResponse
    ) -> str:
        filename = f"{project_name.replace(' ', '_')}_{diff.base_revision}_to_{diff.target_revision}_Revision_Impact.xlsx"
        file_path = os.path.join(REPORT_OUTPUT_DIR, filename)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Revision Impact Analysis"

        headers = ["Signal Tag", "Change Type", "Base Allocation", "Revised Allocation", "Movement Status", "Movement Reason", "Old Properties", "New Properties"]
        cls._apply_header_style(ws, len(headers), f"Revision Diff: {diff.base_revision} → {diff.target_revision}")

        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=2, column=col_idx, value=h)

        add_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
        del_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
        mod_fill = PatternFill(start_color="FEF7E0", end_color="FEF7E0", fill_type="solid")

        for row_idx, item in enumerate(diff.diff_records, 3):
            fill = add_fill if item.change_type == "ADDED" else (del_fill if item.change_type == "DELETED" else (mod_fill if item.change_type == "MODIFIED" else None))
            row_vals = [
                item.tag, item.change_type, item.old_assignment or "—", item.new_assignment or "—",
                item.movement_status, item.movement_reason or "—",
                str(item.old_values) if item.old_values else "—",
                str(item.new_values) if item.new_values else "—"
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                if col_idx == 2 and fill:
                    cell.fill = fill
                    cell.font = Font(name="Calibri", bold=True)
                    cell.alignment = Alignment(horizontal="center")

        cls._auto_fit_columns(ws)
        wb.save(file_path)
        return file_path
