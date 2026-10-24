import os
import sys
import pytest
import openpyxl

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from models import HardwareConfigModel, SignalRecord
from engine.allocator import AllocationEngine
from engine.excel_handler import ExcelHandler
from engine.mapping_engine import MappingEngine
from validation.independent_validator import IndependentValidator

OFFICIAL_MDB = os.path.join(os.path.dirname(__file__), "..", "sample_files", "official", "Customer_IO_MDB.xlsx")
OFFICIAL_REF = os.path.join(os.path.dirname(__file__), "..", "sample_files", "official", "Completed_Allocation_Reference.xlsx")
OFFICIAL_TMPL = os.path.join(os.path.dirname(__file__), "..", "sample_files", "official", "Output_Allocation_Template.xlsx")

def test_official_dataset_reference_comparison():
    assert os.path.exists(OFFICIAL_MDB), f"Customer IO MDB not found at {OFFICIAL_MDB}"
    assert os.path.exists(OFFICIAL_REF), f"Reference allocation not found at {OFFICIAL_REF}"

    # 1. Inspect and parse Customer IO MDB
    inspect_res = ExcelHandler.inspect_workbook(OFFICIAL_MDB)
    sheet_name = inspect_res["sheets"][0]
    headers = inspect_res["headers"][sheet_name]
    column_mapping = inspect_res.get("suggested_mappings") or inspect_res.get("suggested_mapping")

    signals = ExcelHandler.parse_mapped_signals(OFFICIAL_MDB, sheet_name, column_mapping)
    assert len(signals) == 190, f"Expected 190 total records in customer MDB, got {len(signals)}"

    # 2. Run Genuine CP-SAT Allocation Engine with full DeltaV hardware schedule
    config = HardwareConfigModel()
    engine = AllocationEngine(config)
    assignments = engine.run_allocation(signals, full_hardware_schedule=True)

    # 3. Verify Status Breakdown
    assert engine.last_diagnostics is not None
    assert engine.last_diagnostics.solver_name == "Google OR-Tools CP-SAT"
    assert engine.last_diagnostics.is_optimal is True

    # 1,152 physical channels + 1 deleted signal record = 1,153 assignments
    allocated_channels = [a for a in assignments if a.status != "Deleted"]
    assert len(allocated_channels) == 1152, f"Expected exactly 1,152 channels, got {len(allocated_channels)}"

    field_signals = [a for a in allocated_channels if a.tag_status == "Field"]
    alarm_signals = [a for a in allocated_channels if a.tag_status == "Cabinet Alarm"]
    wired_spares = [a for a in allocated_channels if a.tag_status == "Wired Spare"]
    spare_slots = [a for a in allocated_channels if a.tag_status == "Spare Slot"]
    deleted_signals = [a for a in assignments if a.tag_status == "Deleted" or a.status == "Deleted"]

    assert len(field_signals) == 177, f"Expected 177 Field signals, got {len(field_signals)}"
    assert len(alarm_signals) == 12, f"Expected 12 Cabinet Alarms, got {len(alarm_signals)}"
    assert len(wired_spares) == 22, f"Expected 22 Wired Spares, got {len(wired_spares)}"
    assert len(spare_slots) == 941, f"Expected 941 Spare Slots, got {len(spare_slots)}"
    assert len(deleted_signals) == 1, f"Expected 1 Deleted signal, got {len(deleted_signals)}"

    # Verify deleted signal is NOT placed in hardware
    del_sig = deleted_signals[0]
    assert "320-FT" in del_sig.tag and "80002" in del_sig.tag
    assert del_sig.cabinet_id is None or del_sig.cabinet_id == ""

    # 4. Verify Technical Room & Physical Hierarchy Boundaries
    pib30_assigns = [a for a in allocated_channels if a.technical_room == "PIB-30"]
    pib24_assigns = [a for a in allocated_channels if a.technical_room == "PIB-24"]
    assert len(pib30_assigns) == 768  # Cabinets 1 & 2 (8 CIOCs * 96 = 768)
    assert len(pib24_assigns) == 384  # Cabinet 3 (4 CIOCs * 96 = 384)

    unique_cabs = sorted(list(set(a.system_cabinet for a in allocated_channels if a.system_cabinet)))
    assert unique_cabs == ["P24-DCS-50002-01", "P30-DCS-32001-01", "P30-DCS-32001-02"]

    unique_ctrls = sorted(list(set(a.controller for a in allocated_channels if a.controller)))
    assert unique_ctrls == ["P24DPK500201", "P30DPK320101", "P30DPK320102"]

    unique_ciocs = sorted(list(set(a.cioc_id for a in allocated_channels if a.cioc_id)))
    assert len(unique_ciocs) == 12

    # 5. Independent Validator must pass with 0 hard errors
    val_config = HardwareConfigModel(cabinets_count=3, ciocs_per_cabinet=4, baseplates_per_cioc=8, charms_per_baseplate=12)
    val_config.rules.enforce_ab_redundancy_separation = True
    validator = IndependentValidator(val_config)
    summary = validator.validate(signals, assignments)

    assert summary.error_count == 0, f"Independent validator found {summary.error_count} errors: {[i.message for i in summary.issues if i.severity == 'ERROR']}"

    # 6. Template-Driven Excel Export Verification
    out_file = ExcelHandler.export_completed_allocation("Test_Reference_Verification", "Rev_00", assignments, signals)
    assert os.path.exists(out_file)

    wb = openpyxl.load_workbook(out_file, data_only=True)
    ws = wb["IO Assignment"] if "IO Assignment" in wb.sheetnames else wb.active
    assert ws.max_column == 76
    assert ws.max_row == 1153  # Row 1 headers + 1,152 channels
    row2_vals = [c.value for c in ws[2]]
    assert row2_vals[0] == 1  # Sr. No. 1
    assert row2_vals[2] is not None  # Tag Name
