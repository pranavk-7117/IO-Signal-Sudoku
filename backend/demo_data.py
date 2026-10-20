import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from typing import List, Dict, Any
from database import get_db
from models import (
    ProjectCreate, RevisionCreate, HardwareConfigModel, RuleConfigModel, SignalRecord
)
from db_service import DBService
from engine.allocator import AllocationEngine
from engine.validator import ValidationEngine
from engine.tq_generator import TQGenerator

ORIG_SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "sample_files")
if os.environ.get("VERCEL"):
    SAMPLE_DIR = "/tmp/sample_files"
    try:
        os.makedirs(SAMPLE_DIR, exist_ok=True)
        if os.path.exists(ORIG_SAMPLE_DIR):
            import shutil
            for f in os.listdir(ORIG_SAMPLE_DIR):
                src_f = os.path.join(ORIG_SAMPLE_DIR, f)
                dst_f = os.path.join(SAMPLE_DIR, f)
                if not os.path.exists(dst_f):
                    shutil.copy2(src_f, dst_f)
    except Exception:
        pass
else:
    SAMPLE_DIR = ORIG_SAMPLE_DIR
    try:
        os.makedirs(SAMPLE_DIR, exist_ok=True)
    except Exception:
        SAMPLE_DIR = "/tmp/sample_files"
        os.makedirs(SAMPLE_DIR, exist_ok=True)

def create_sample_excel_files():
    """Generates synthetic customer IO workbooks and output template."""
    reva_path = os.path.join(SAMPLE_DIR, "Customer_IO_List_RevA.xlsx")
    revb_path = os.path.join(SAMPLE_DIR, "Customer_IO_List_RevB.xlsx")
    template_path = os.path.join(SAMPLE_DIR, "DeltaV_Allocation_Template.xlsx")

    # Generate Rev A IO Workbook
    wb_a = openpyxl.Workbook()
    ws_a = wb_a.active
    ws_a.title = "Master_IO_Database"

    headers = ["Tag", "Signal Description", "IO Type", "Field Cable Tag", "Wire Pair", "Redundancy Designator", "Process Unit / Logical Group", "Segregation"]
    ws_a.append(headers)

    signals_data_a = [
        # Boiler Feedwater System (AI, AO, DI, DO) - Cable CB-01 (6 pairs)
        ("PT-1001", "Feedwater Pump Discharge Pressure", "AI", "CB-01", 1, "A", "Boiler Feedwater", "DCS"),
        ("PT-1002", "Feedwater Pump Suction Pressure", "AI", "CB-01", 2, "A", "Boiler Feedwater", "DCS"),
        ("FT-1001", "Main Boiler Feedwater Mass Flow", "AI", "CB-01", 3, "A", "Boiler Feedwater", "DCS"),
        ("TT-1001", "Economizer Inlet Temperature", "AI", "CB-01", 4, "A", "Boiler Feedwater", "DCS"),
        ("LT-1001", "Deaerator Storage Tank Level", "AI", "CB-01", 5, "A", "Boiler Feedwater", "DCS"),
        ("FV-1001", "Feedwater Control Valve Positioner Feedback", "AO", "CB-01", 6, "A", "Boiler Feedwater", "DCS"),

        # Redundant B Channel for Boiler Feedwater - Cable CB-02 (6 pairs) - Must land on separate CIOC!
        ("PT-1001B", "Feedwater Pump Discharge Pressure [Redundant B]", "AI", "CB-02", 1, "B", "Boiler Feedwater", "DCS"),
        ("PT-1002B", "Feedwater Pump Suction Pressure [Redundant B]", "AI", "CB-02", 2, "B", "Boiler Feedwater", "DCS"),
        ("FT-1001B", "Main Boiler Feedwater Mass Flow [Redundant B]", "AI", "CB-02", 3, "B", "Boiler Feedwater", "DCS"),
        ("TT-1001B", "Economizer Inlet Temperature [Redundant B]", "AI", "CB-02", 4, "B", "Boiler Feedwater", "DCS"),
        ("LT-1001B", "Deaerator Storage Tank Level [Redundant B]", "AI", "CB-02", 5, "B", "Boiler Feedwater", "DCS"),
        ("FV-1001B", "Feedwater Control Valve Secondary Loop AO", "AO", "CB-02", 6, "B", "Boiler Feedwater", "DCS"),

        # Cooling Water Circuit - Cable CB-03 (4 pairs)
        ("FT-2001", "Cooling Tower Circulation Flow", "AI", "CB-03", 1, "", "Cooling Water", "DCS"),
        ("TT-2001", "Cooling Water Supply Temp", "AI", "CB-03", 2, "", "Cooling Water", "DCS"),
        ("TT-2002", "Cooling Water Return Temp", "AI", "CB-03", 3, "", "Cooling Water", "DCS"),
        ("PT-2001", "Cooling Water Header Pressure", "AI", "CB-03", 4, "", "Cooling Water", "DCS"),

        # Steam Turbine Generator - Discrete IO - Cable CB-04 (6 pairs)
        ("HS-3001", "Turbine Trip Pushbutton Status", "DI", "CB-04", 1, "A", "Turbine Controls", "ESD"),
        ("ZS-3001", "Main Steam Stop Valve Fully Open", "DI", "CB-04", 2, "A", "Turbine Controls", "ESD"),
        ("ZS-3002", "Main Steam Stop Valve Fully Closed", "DI", "CB-04", 3, "A", "Turbine Controls", "ESD"),
        ("XV-3001", "Turbine Emergency Trip Solenoid", "DO-SOV", "CB-04", 4, "A", "Turbine Controls", "ESD"),
        ("XV-3002", "Vacuum Breaker Solenoid Valve", "DO-SOV", "CB-04", 5, "A", "Turbine Controls", "ESD"),
        ("ALM-3001", "Turbine Overspeed Alarm Strobe", "DO", "CB-04", 6, "A", "Turbine Controls", "DCS"),

        # Redundant Turbine Controls B - Cable CB-05 (6 pairs)
        ("HS-3001B", "Turbine Trip Pushbutton Status [Redundant B]", "DI", "CB-05", 1, "B", "Turbine Controls", "ESD"),
        ("ZS-3001B", "Main Steam Stop Valve Open [Redundant B]", "DI", "CB-05", 2, "B", "Turbine Controls", "ESD"),
        ("ZS-3002B", "Main Steam Stop Valve Closed [Redundant B]", "DI", "CB-05", 3, "B", "Turbine Controls", "ESD"),
        ("XV-3001B", "Turbine Emergency Trip Solenoid B", "DO-SOV", "CB-05", 4, "B", "Turbine Controls", "ESD"),
        ("XV-3002B", "Vacuum Breaker Solenoid Valve B", "DO-SOV", "CB-05", 5, "B", "Turbine Controls", "ESD"),
        ("ALM-3001B", "Turbine Overspeed Alarm Horn", "DO", "CB-05", 6, "B", "Turbine Controls", "DCS"),

        # Distillation Column Unit - Cable CB-06 (8 pairs)
        ("PT-4001", "Column Overhead Vapor Pressure", "AI", "CB-06", 1, "", "Distillation Unit", "DCS"),
        ("TT-4001", "Column Tray 12 Temperature", "AI", "CB-06", 2, "", "Distillation Unit", "DCS"),
        ("TT-4002", "Column Reboiler Outlet Temp", "AI", "CB-06", 3, "", "Distillation Unit", "DCS"),
        ("LT-4001", "Reflux Drum Hydrocarbon Level", "AI", "CB-06", 4, "", "Distillation Unit", "DCS"),
        ("LT-4002", "Column Sump Liquid Level", "AI", "CB-06", 5, "", "Distillation Unit", "DCS"),
        ("FV-4001", "Reflux Return Flow Control Valve", "AO", "CB-06", 6, "", "Distillation Unit", "DCS"),
        ("FV-4002", "Reboiler Steam Flow Control Valve", "AO", "CB-06", 7, "", "Distillation Unit", "DCS"),
        ("PV-4001", "Offgas Vent Pressure Control Valve", "AO", "CB-06", 8, "", "Distillation Unit", "DCS"),

        # Synthetic Deliberate Issues for demonstration
        ("TAG-ERR-901", "Special Pyrometer Pyrometric Sensor", "UNKNOWN_OPTIC_440V", "CB-99", 1, "", "Special Instrumentation", "DCS"),
    ]

    for r in signals_data_a:
        ws_a.append(r)

    wb_a.save(reva_path)

    # Generate Rev B IO Workbook (Modified, Added, Deleted signals)
    wb_b = openpyxl.Workbook()
    ws_b = wb_b.active
    ws_b.title = "Master_IO_Database"
    ws_b.append(headers)

    signals_data_b = [
        # Same boiler signals (should PRESERVE allocation)
        ("PT-1001", "Feedwater Pump Discharge Pressure", "AI", "CB-01", 1, "A", "Boiler Feedwater", "DCS"),
        ("PT-1002", "Feedwater Pump Suction Pressure (Range Adjusted)", "AI", "CB-01", 2, "A", "Boiler Feedwater", "DCS"), # MODIFIED desc
        ("FT-1001", "Main Boiler Feedwater Mass Flow", "AI", "CB-01", 3, "A", "Boiler Feedwater", "DCS"),
        ("TT-1001", "Economizer Inlet Temperature", "AI", "CB-01", 4, "A", "Boiler Feedwater", "DCS"),
        ("LT-1001", "Deaerator Storage Tank Level", "AI", "CB-01", 5, "A", "Boiler Feedwater", "DCS"),
        ("FV-1001", "Feedwater Control Valve Positioner Feedback", "AO", "CB-01", 6, "A", "Boiler Feedwater", "DCS"),

        # Redundant B Channel
        ("PT-1001B", "Feedwater Pump Discharge Pressure [Redundant B]", "AI", "CB-02", 1, "B", "Boiler Feedwater", "DCS"),
        ("PT-1002B", "Feedwater Pump Suction Pressure [Redundant B]", "AI", "CB-02", 2, "B", "Boiler Feedwater", "DCS"),
        ("FT-1001B", "Main Boiler Feedwater Mass Flow [Redundant B]", "AI", "CB-02", 3, "B", "Boiler Feedwater", "DCS"),
        ("TT-1001B", "Economizer Inlet Temperature [Redundant B]", "AI", "CB-02", 4, "B", "Boiler Feedwater", "DCS"),
        ("LT-1001B", "Deaerator Storage Tank Level [Redundant B]", "AI", "CB-02", 5, "B", "Boiler Feedwater", "DCS"),
        ("FV-1001B", "Feedwater Control Valve Secondary Loop AO", "AO", "CB-02", 6, "B", "Boiler Feedwater", "DCS"),

        # Cooling Water - PT-2001 was deleted, TT-2003 was ADDED
        ("FT-2001", "Cooling Tower Circulation Flow", "AI", "CB-03", 1, "", "Cooling Water", "DCS"),
        ("TT-2001", "Cooling Water Supply Temp", "AI", "CB-03", 2, "", "Cooling Water", "DCS"),
        ("TT-2002", "Cooling Water Return Temp", "AI", "CB-03", 3, "", "Cooling Water", "DCS"),
        ("TT-2003", "Cooling Tower Basin Makeup Temp [NEW]", "AI", "CB-03", 4, "", "Cooling Water", "DCS"), # ADDED

        # Turbine discrete signals
        ("HS-3001", "Turbine Trip Pushbutton Status", "DI", "CB-04", 1, "A", "Turbine Controls", "ESD"),
        ("ZS-3001", "Main Steam Stop Valve Fully Open", "DI", "CB-04", 2, "A", "Turbine Controls", "ESD"),
        ("ZS-3002", "Main Steam Stop Valve Fully Closed", "DI", "CB-04", 3, "A", "Turbine Controls", "ESD"),
        ("XV-3001", "Turbine Emergency Trip Solenoid", "DO-SOV", "CB-04", 4, "A", "Turbine Controls", "ESD"),
        ("XV-3002", "Vacuum Breaker Solenoid Valve", "DO-SOV", "CB-04", 5, "A", "Turbine Controls", "ESD"),
        ("ALM-3001", "Turbine Overspeed Alarm Strobe", "DO", "CB-04", 6, "A", "Turbine Controls", "DCS"),

        # Redundant Turbine B
        ("HS-3001B", "Turbine Trip Pushbutton Status [Redundant B]", "DI", "CB-05", 1, "B", "Turbine Controls", "ESD"),
        ("ZS-3001B", "Main Steam Stop Valve Open [Redundant B]", "DI", "CB-05", 2, "B", "Turbine Controls", "ESD"),
        ("ZS-3002B", "Main Steam Stop Valve Closed [Redundant B]", "DI", "CB-05", 3, "B", "Turbine Controls", "ESD"),
        ("XV-3001B", "Turbine Emergency Trip Solenoid B", "DO-SOV", "CB-05", 4, "B", "Turbine Controls", "ESD"),
        ("XV-3002B", "Vacuum Breaker Solenoid Valve B", "DO-SOV", "CB-05", 5, "B", "Turbine Controls", "ESD"),
        ("ALM-3001B", "Turbine Overspeed Alarm Horn", "DO", "CB-05", 6, "B", "Turbine Controls", "DCS"),

        # Distillation Column Unit
        ("PT-4001", "Column Overhead Vapor Pressure", "AI", "CB-06", 1, "", "Distillation Unit", "DCS"),
        ("TT-4001", "Column Tray 12 Temperature", "AI", "CB-06", 2, "", "Distillation Unit", "DCS"),
        ("TT-4002", "Column Reboiler Outlet Temp", "AI", "CB-06", 3, "", "Distillation Unit", "DCS"),
        ("LT-4001", "Reflux Drum Hydrocarbon Level", "AI", "CB-06", 4, "", "Distillation Unit", "DCS"),
        ("LT-4002", "Column Sump Liquid Level", "AI", "CB-06", 5, "", "Distillation Unit", "DCS"),
        ("FV-4001", "Reflux Return Flow Control Valve", "AO", "CB-06", 6, "", "Distillation Unit", "DCS"),
        ("FV-4002", "Reboiler Steam Flow Control Valve", "AO", "CB-06", 7, "", "Distillation Unit", "DCS"),
        ("PV-4001", "Offgas Vent Pressure Control Valve", "AO", "CB-06", 8, "", "Distillation Unit", "DCS"),
        ("AT-4001", "Overhead Analyzer Chromatograph Purity", "AI", "CB-06", 9, "", "Distillation Unit", "DCS"), # ADDED
    ]

    for r in signals_data_b:
        ws_b.append(r)

    wb_b.save(revb_path)

    # Generate Output Allocation Template
    wb_t = openpyxl.Workbook()
    ws_t = wb_t.active
    ws_t.title = "Allocated_Hardware_Schedule"
    template_headers = [
        "Tag Name", "Signal Description", "Signal Type", "Cable Tag", "Pair #",
        "Redundancy", "Process Unit", "Assigned Cabinet", "Assigned CIOC",
        "Assigned Baseplate", "CHARM Slot", "CHARM Hardware Type", "Terminal Block Type", "Allocation Status"
    ]
    ws_t.append(template_headers)
    wb_t.save(template_path)

    return reva_path, revb_path, template_path

def seed_demo_project():
    """Seeds the demonstration project into SQLite with synthetic data, runs allocation, independent validator & TQ generation."""
    reva_path, revb_path, template_path = create_sample_excel_files()

    conn = get_db()
    cursor = conn.cursor()
    # Check if Demo project exists
    cursor.execute("SELECT id FROM projects WHERE project_id = 'PRJ-DEMO-DELTAV';")
    existing = cursor.fetchone()
    if existing:
        conn.close()
        return existing["id"]

    conn.close()

    # Create project
    proj_data = ProjectCreate(
        project_id="PRJ-DEMO-DELTAV",
        name="Demo DeltaV Allocation",
        customer_name="PetroChem Refining Corporation",
        description="Comprehensive synthetic DCS IO Allocation demonstration project for DeltaV CHARMs electronic marshalling.",
        system="DeltaV v14.3 / Electronic Marshalling CHARMs"
    )
    proj = DBService.create_project(proj_data)
    p_id = proj["id"]

    # Get Rev A
    revs = DBService.get_revisions(p_id)
    rev_a_id = revs[0]["id"]

    # Parse and insert Rev A signals
    mapping = {
        "tag": "Tag",
        "description": "Signal Description",
        "signal_type": "IO Type",
        "cable_id": "Field Cable Tag",
        "pair_number": "Wire Pair",
        "redundancy": "Redundancy Designator",
        "logical_group": "Process Unit / Logical Group"
    }
    from engine.excel_handler import ExcelHandler
    signals = ExcelHandler.parse_mapped_signals(reva_path, "Master_IO_Database", mapping)
    DBService.save_imported_signals(p_id, rev_a_id, signals, "Customer_IO_List_RevA.xlsx")

    # Run allocation engine
    hw_config = DBService.get_hardware_config(p_id, rev_a_id)
    engine = AllocationEngine(hw_config)
    assignments = engine.run_allocation(signals)
    DBService.save_assignments(p_id, rev_a_id, assignments)

    # Run independent validation engine
    validator = ValidationEngine(hw_config)
    validation_summary = validator.validate(signals, assignments)
    DBService.save_validation_results(p_id, rev_a_id, validation_summary.issues)

    # Generate TQs
    tqs = TQGenerator.generate_tqs_from_validation(p_id, rev_a_id, validation_summary.issues, assignments)
    for tq in tqs:
        DBService.create_or_update_tq(tq)

    # Also pre-create Rev B for immediate revision comparison demonstrations
    rev_b_data = RevisionCreate(
        revision_code="Rev B",
        description="Customer Revised IO List (Deleted PT-2001, Added TT-2003 and AT-4001, modified ranges)"
    )
    rev_b = DBService.create_revision(p_id, rev_b_data)
    rev_b_id = rev_b["id"]

    signals_b = ExcelHandler.parse_mapped_signals(revb_path, "Master_IO_Database", mapping)
    DBService.save_imported_signals(p_id, rev_b_id, signals_b, "Customer_IO_List_RevB.xlsx")

    # Run allocation with preservation from Rev A
    prev_assignments_map = {a.tag: a for a in assignments}
    assignments_b = engine.run_allocation(signals_b, previous_assignments=prev_assignments_map)
    DBService.save_assignments(p_id, rev_b_id, assignments_b)

    val_b = validator.validate(signals_b, assignments_b)
    DBService.save_validation_results(p_id, rev_b_id, val_b.issues)

    # Switch active revision back to Rev A initially so engineer walks through the full journey
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute("UPDATE projects SET current_revision = 'Rev A' WHERE id = ?;", (p_id,))
    cursor.execute("UPDATE revisions SET is_active = 0 WHERE project_id = ?;", (p_id,))
    cursor.execute("UPDATE revisions SET is_active = 1 WHERE id = ?;", (rev_a_id,))
    conn.commit()
    conn.close()

    return p_id
